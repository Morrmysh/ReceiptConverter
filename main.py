import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, PageBreak, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os
import datetime
import sys
import zipfile
import io
from typing import Any

pdfmetrics.registerFont(TTFont('TimesNewRoman', 'assets/times.ttf'))
pdfmetrics.registerFont(TTFont('TimesNewRoman-Bold', 'assets/timesbd.ttf'))

def receipt_to_pdf(excel_file_path):
    try:
        wb = load_workbook(excel_file_path)
        sheet = wb.active

        current_date = datetime.datetime.now().strftime('%Y-%m-%d')

        output_folder = os.path.join(os.path.dirname(excel_file_path), "Квитанции от " + current_date)
        os.makedirs(output_folder, exist_ok=True)
        pdf_file_path = os.path.join(str(output_folder), str(os.path.basename(excel_file_path).replace('.xlsx', '.pdf')))

        pdf = SimpleDocTemplate(str(pdf_file_path), pagesize=landscape(A4), title=str(os.path.basename(excel_file_path).replace('.xlsx', '')))
        pdf.topMargin = 10
        pdf.bottomMargin = 10
        pdf.leftMargin = 20
        pdf.rightMargin = 20

        header_style = ParagraphStyle(
            name='HeaderStyle',
            fontName='TimesNewRoman',
            fontSize=14,
            spaceAfter=12,
        )
        text_style = ParagraphStyle(
            name='TextStyle',
            fontName='TimesNewRoman',
            fontSize=12,
            spaceAfter=10,
            spaceBefore=10,
        )
        cell_style = ParagraphStyle(
            name='CellStyle',
            fontName='TimesNewRoman',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=0,
        )
        cell_style_bold = ParagraphStyle(
            name='CellStyleBold',
            fontName='TimesNewRoman-Bold',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=0,
        )
        center_cell_style = ParagraphStyle(
            name='CenterCellStyle',
            fontName='TimesNewRoman',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=1,
        )
        center_cell_style_bold = ParagraphStyle(
            name='CenterCellStyleBold',
            fontName='TimesNewRoman-Bold',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=1,
        )

        content = []

        title = Paragraph(
            f"Квитанция номер {sheet['B4'].value.strip()}.",
            header_style,
        )
        content.append(title)

        excel_data = [[str(cell.value) for cell in row] for row in sheet['A5:D8']]
        excel_data.insert(1, ["ИНН:", "5042043326", "", ""])
        data = []
        for row in excel_data:
            data.append([])
            for num, cell in enumerate(row):
                if num % 2 == 0:
                    data[-1].append(Paragraph(cell, cell_style_bold))
                else:
                    data[-1].append(Paragraph(cell, cell_style))

        info_table = Table(data, style=[
            ('SPAN', (2, 0), (2, 1)),
            ('SPAN', (3, 0), (3, 1)),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOX', (0, 0), (-1, -1), 1, colors.black),
            ('LINEABOVE', (0, 0), (-1, -1), 1, colors.black),
            ('LINEAFTER', (1, 0), (1, -1), 1, colors.black)
        ],
                           colWidths=[120, None, 140, None],
                           hAlign='LEFT')
        content.append(info_table)

        text = Paragraph("Расшифровка по видам услуг", text_style)
        content.append(text)

        excel_data = [[str(cell.value if cell.value is not None else '').replace('_x000D_\n', '') for cell in row] for row in sheet['A12:P25']]
        data = []
        for row in excel_data[:3]:
            data.append([])
            for cell in row:
                data[-1].append(Paragraph(cell, center_cell_style_bold))
        for row in excel_data[3:]:
            data.append([])
            data[-1].append(Paragraph(row[0], cell_style_bold))
            for cell in row[1:]:
                data[-1].append(Paragraph(cell, center_cell_style))

        summary_table = Table(data, style=[
            ('SPAN', (0, 0), (0, 2)),
            ('SPAN', (1, 0), (1, 2)),
            ('SPAN', (2, 0), (2, 2)),
            ('SPAN', (3, 0), (3, 2)),
            ('SPAN', (4, 0), (4, 2)),
            ('SPAN', (5, 0), (8, 0)),
            ('SPAN', (5, 1), (6, 1)),
            ('SPAN', (7, 1), (8, 1)),
            ('SPAN', (5, 2), (7, 2)),
            ('SPAN', (9, 0), (9, 2)),
            ('SPAN', (10, 0), (11, 0)),
            ('SPAN', (10, 1), (10, 2)),
            ('SPAN', (11, 1), (11, 2)),
            ('SPAN', (12, 0), (13, 0)),
            ('SPAN', (12, 1), (12, 2)),
            ('SPAN', (13, 1), (13, 2)),
            ('SPAN', (14, 0), (14, 2)),
            ('SPAN', (15, 0), (15, 2)),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BACKGROUND', (0, 0), (-1, 2), colors.lightgrey),
            ('BACKGROUND', (0, 3), (-1, 4), colors.whitesmoke),
            ('BACKGROUND', (0, 6), (-1, 6), colors.whitesmoke),
            ('BACKGROUND', (0, 12), (-1, 12), colors.whitesmoke),
            ('BACKGROUND', (0, 3), (0, -1), colors.whitesmoke),
        ],
                              colWidths=[100, 60, 60, 50, 40, 30, 55, 30, 40, 40, 40, 40, 60, None, None, None],)
        content.append(summary_table)

        data = [[Paragraph('Мясникова Ангелина Михайловна', text_style),
                Paragraph(datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), text_style)]]
        table = Table(data, style=[('TOPPADDING', (0, 0), (-1, -1), 10), ('LEFTPADDING', (0, 0), (-1, -1), 0)])
        content.append(table)

        pdf.build(content)

        return pdf_file_path

    except Exception as e:
        error_msg = str(e)
        print(f"Error in excel_to_pdf: {error_msg}", file=sys.stderr)
        messagebox.showerror("Error", error_msg)
        return None
    finally:
        if 'wb' in locals():
            wb.close()


def open_receipt():
    file_paths = filedialog.askopenfilename(filetypes=[('Excel files', '*.xlsx')], multiple=True)
    if file_paths:
        # Создаем окно загрузки
        loading_window = tk.Toplevel(root)
        loading_window.title("Обработка...")
        loading_window.geometry("300x100")
        loading_window.resizable(False, False)
        
        # Центрируем окно по экрану
        center_window(loading_window)
        loading_window.transient(root)
        loading_window.grab_set()
        
        label = tk.Label(loading_window, text="Обработка Excel-файлов...", font=('Arial', 12))
        label.pack(expand=True)
        
        # Обновляем GUI перед обработкой
        root.update()
        
        success_count = 0
        error_count = 0
        for file_path in file_paths:
            pdf_file_path = receipt_to_pdf(file_path)
            if pdf_file_path:
                success_count += 1
            else:
                error_count += 1
        
        # Закрываем окно загрузки
        loading_window.destroy()
        
        if success_count > 0:
            messagebox.showinfo("Success", f"Создано квитанций: {success_count}")
        if error_count > 0:
            messagebox.showwarning("Errors", f"Не удалось создать: {error_count} квитанций")


def overhaul_to_pdf(excel_file_path):
    try:
        wb = load_workbook(excel_file_path)
        sheet = wb.active

        # Извлекаем QR-код из Excel через zipfile в память
        qr_image = None
        with zipfile.ZipFile(excel_file_path, 'r') as zip_ref:
            media_files = [f for f in zip_ref.namelist() if f.startswith('xl/media/')]
            if media_files:
                qr_bytes = io.BytesIO()
                with zip_ref.open(media_files[0]) as source:
                    qr_bytes.write(source.read())
                qr_bytes.seek(0)
                qr_image = Image(qr_bytes, width=152, height=152)

        current_date = datetime.datetime.now().strftime('%Y-%m-%d')
        output_folder = os.path.join(os.path.dirname(excel_file_path), "Квитанции капремонт от " + current_date)
        os.makedirs(output_folder, exist_ok=True)
        pdf_file_path = os.path.join(str(output_folder), str(os.path.basename(excel_file_path).replace('.xlsx', '.pdf')))

        pdf = SimpleDocTemplate(str(pdf_file_path), pagesize=landscape(A4), title=str(os.path.basename(excel_file_path).replace('.xlsx', '')))
        pdf.topMargin = 10
        pdf.bottomMargin = 10
        pdf.leftMargin = 20
        pdf.rightMargin = 20

        cell_style = ParagraphStyle(
            name='CellStyle',
            fontName='TimesNewRoman',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=0,
        )
        cell_style_bold = ParagraphStyle(
            name='CellStyleBold',
            fontName='TimesNewRoman-Bold',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=0,
        )
        cell_style_bold_12 = ParagraphStyle(
            name='CellStyleBold',
            fontName='TimesNewRoman-Bold',
            fontSize=12,
            spaceAfter=10,
            wordWrap=None,
            alignment=0,
        )
        center_cell_style = ParagraphStyle(
            name='CenterCellStyle',
            fontName='TimesNewRoman',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=1,
        )
        center_cell_style_bold = ParagraphStyle(
            name='CenterCellStyleBold',
            fontName='TimesNewRoman-Bold',
            fontSize=10,
            spaceAfter=10,
            wordWrap=None,
            alignment=1,
        )

        content = []

        for cursor in range(1, sheet.max_row + 1, 32):

            overhaul_data = []
            end = cursor + 31
            for row in range(cursor, end + 1):
                row_data = []
                for col in range(1, 56):
                    cell_value = sheet.cell(row=row, column=col).value
                    row_data.append(cell_value if cell_value is not None else '')
                overhaul_data.append(row_data)

            table_data: list[list[Any]] = [[None for _ in range(17)] for _ in range(19)]

            table_data[0][0] = Paragraph(str(overhaul_data[0][0]), center_cell_style_bold)

            table_data[1][0] = Paragraph(str(overhaul_data[2][0]), cell_style_bold_12)
            table_data[1][2] = Paragraph(str(overhaul_data[2][2]), cell_style_bold_12)
            table_data[1][8] = Paragraph(str(overhaul_data[2][25]), cell_style_bold_12)
            table_data[1][9] = Paragraph(str(overhaul_data[2][27]), cell_style_bold_12)

            table_data[2][0] = Paragraph(str(overhaul_data[3][0]), cell_style_bold)
            table_data[2][2] = Paragraph(str(overhaul_data[3][1]) + " " + str(overhaul_data[3][6]), cell_style)
            table_data[2][8] = Paragraph(str(overhaul_data[3][25]), cell_style_bold)
            table_data[2][9] = Paragraph(str(overhaul_data[3][32]), cell_style_bold)
            table_data[2][12] = Paragraph(str(overhaul_data[3][42]), cell_style_bold)
            table_data[2][13] = Paragraph(str(overhaul_data[3][46]), cell_style_bold)
            table_data[2][14] = Paragraph(str(overhaul_data[3][52]), cell_style_bold)

            table_data[3][0] = Paragraph(str(overhaul_data[4][0]), cell_style_bold)
            table_data[3][8] = Paragraph(str(overhaul_data[6][25]), cell_style)
            table_data[3][9] = Paragraph(str(overhaul_data[6][32]), cell_style)
            table_data[3][12] = Paragraph(str(overhaul_data[6][42]), cell_style)
            table_data[3][13] = Paragraph(str(overhaul_data[6][46]), cell_style)
            table_data[3][14] = Paragraph(str(overhaul_data[6][52]), cell_style)

            table_data[4][0] = Paragraph(str(overhaul_data[5][0]), cell_style)
            table_data[4][8] = Paragraph(str(overhaul_data[9][25]), center_cell_style_bold)
            table_data[4][12] = qr_image if qr_image else Paragraph("QR-code", center_cell_style_bold)

            table_data[5][0] = Paragraph(str(overhaul_data[6][0]), cell_style_bold)
            table_data[5][2] = Paragraph(str(overhaul_data[6][2]), cell_style)
            table_data[5][8] = Paragraph(str(overhaul_data[10][25]), cell_style_bold)
            table_data[5][11] = Paragraph(str(overhaul_data[10][32]), center_cell_style)

            table_data[6][0] = Paragraph(str(overhaul_data[7][0]), cell_style_bold)
            table_data[6][2] = Paragraph(str(overhaul_data[7][6]) + " " + str(overhaul_data[7][11]), cell_style)
            table_data[6][4] = Paragraph(str(overhaul_data[7][15]), cell_style_bold)
            table_data[6][6] = Paragraph(str(overhaul_data[7][21]) + " " + str(overhaul_data[7][23]), cell_style)
            table_data[6][8] = Paragraph(str(overhaul_data[11][25]), cell_style_bold)
            table_data[6][11] = Paragraph(str(overhaul_data[11][32]), center_cell_style)

            table_data[7][0] = Paragraph(str(overhaul_data[8][0]), cell_style_bold)
            table_data[7][2] = Paragraph(str(overhaul_data[8][7]), cell_style)
            table_data[7][4] = Paragraph(str(overhaul_data[8][19]), cell_style_bold)
            table_data[7][6] = Paragraph(str(overhaul_data[8][20]), cell_style)
            table_data[7][8] = Paragraph(str(overhaul_data[12][25]), cell_style_bold)
            table_data[7][11] = Paragraph(str(overhaul_data[12][32]), center_cell_style)

            table_data[8][0] = Paragraph(str(overhaul_data[9][0]), cell_style_bold)
            table_data[8][2] = Paragraph(str(overhaul_data[9][2]), cell_style)
            table_data[8][8] = Paragraph(str(overhaul_data[13][25]), cell_style_bold)
            table_data[8][11] = Paragraph(str(overhaul_data[13][32]), center_cell_style)

            table_data[9][8] = Paragraph(str(overhaul_data[19][25]), cell_style_bold)
            table_data[9][11] = Paragraph(str(overhaul_data[19][32]), center_cell_style_bold)

            table_data[11][0] = Paragraph(str(overhaul_data[21][0]), cell_style_bold_12)
            table_data[11][2] = Paragraph(str(overhaul_data[21][2]), cell_style_bold_12)

            table_data[12][0] = Paragraph(str(overhaul_data[22][0]), cell_style_bold)
            table_data[12][1] = Paragraph(str(overhaul_data[22][7]), cell_style_bold)
            table_data[12][2] = Paragraph(str(overhaul_data[22][9]), cell_style_bold)
            table_data[12][6] = Paragraph(str(overhaul_data[22][15]), cell_style_bold)
            table_data[12][7] = Paragraph(str(overhaul_data[22][17]), cell_style_bold)
            table_data[12][9] = Paragraph(str(overhaul_data[22][21]), cell_style_bold)
            table_data[12][10] = Paragraph(str(overhaul_data[22][24]), cell_style_bold)
            table_data[12][11] = Paragraph(str(overhaul_data[22][27]), cell_style_bold)
            table_data[12][12] = Paragraph(str(overhaul_data[22][30]), cell_style_bold)
            table_data[12][13] = Paragraph(str(overhaul_data[22][32]), cell_style_bold)
            table_data[12][14] = Paragraph(str(overhaul_data[22][34]), cell_style_bold)

            table_data[13][14] = Paragraph(str(overhaul_data[24][34]), cell_style_bold)
            table_data[13][15] = Paragraph(str(overhaul_data[24][36]), cell_style_bold)

            table_data[14][2] = Paragraph(str(overhaul_data[25][9]), cell_style_bold)
            table_data[14][3] = Paragraph(str(overhaul_data[25][11]), cell_style_bold)
            table_data[14][4] = Paragraph(str(overhaul_data[25][12]), cell_style_bold)
            table_data[14][5] = Paragraph(str(overhaul_data[25][14]), cell_style_bold)
            table_data[14][7] = Paragraph(str(overhaul_data[25][17]), cell_style_bold)
            table_data[14][8] = Paragraph(str(overhaul_data[25][19]), cell_style_bold)
            table_data[14][15] = Paragraph(str(overhaul_data[25][36]), cell_style_bold)
            table_data[14][16] = Paragraph(str(overhaul_data[25][38]), cell_style_bold)

            table_data[15][0] = Paragraph(str(overhaul_data[27][0]), center_cell_style_bold)
            table_data[15][1] = Paragraph(str(overhaul_data[27][7]), center_cell_style_bold)
            table_data[15][2] = Paragraph(str(overhaul_data[27][9]), center_cell_style_bold)
            table_data[15][4] = Paragraph(str(overhaul_data[27][12]), center_cell_style_bold)
            table_data[15][6] = Paragraph(str(overhaul_data[27][15]), center_cell_style_bold)
            table_data[15][7] = Paragraph(str(overhaul_data[27][17]), center_cell_style_bold)
            table_data[15][8] = Paragraph(str(overhaul_data[27][19]), center_cell_style_bold)
            table_data[15][9] = Paragraph(str(overhaul_data[27][21]), center_cell_style_bold)
            table_data[15][10] = Paragraph(str(overhaul_data[27][24]), center_cell_style_bold)
            table_data[15][11] = Paragraph(str(overhaul_data[27][27]), center_cell_style_bold)
            table_data[15][12] = Paragraph(str(overhaul_data[27][30]), center_cell_style_bold)
            table_data[15][13] = Paragraph(str(overhaul_data[27][32]), center_cell_style_bold)
            table_data[15][14] = Paragraph(str(overhaul_data[27][34]), center_cell_style_bold)
            table_data[15][15] = Paragraph(str(overhaul_data[27][36]), center_cell_style_bold)
            table_data[15][16] = Paragraph(str(overhaul_data[27][38]), center_cell_style_bold)

            table_data[16][0] = Paragraph(str(overhaul_data[28][0]), cell_style)
            table_data[16][1] = Paragraph(str(overhaul_data[28][7]), center_cell_style)
            table_data[16][2] = Paragraph(str(overhaul_data[28][9]), center_cell_style)
            table_data[16][3] = Paragraph(str(overhaul_data[28][11]), center_cell_style)
            table_data[16][4] = Paragraph(str(overhaul_data[28][12]), center_cell_style)
            table_data[16][5] = Paragraph(str(overhaul_data[28][14]), center_cell_style)
            table_data[16][6] = Paragraph(str(overhaul_data[28][15]), center_cell_style)
            table_data[16][7] = Paragraph(str(overhaul_data[28][17]), center_cell_style)
            table_data[16][8] = Paragraph(str(overhaul_data[28][19]), center_cell_style)
            table_data[16][9] = Paragraph(str(overhaul_data[28][21]), center_cell_style)
            table_data[16][10] = Paragraph(str(overhaul_data[28][24]), center_cell_style)
            table_data[16][11] = Paragraph(str(overhaul_data[28][27]), center_cell_style)
            table_data[16][12] = Paragraph(str(overhaul_data[28][30]), center_cell_style)
            table_data[16][13] = Paragraph(str(overhaul_data[28][32]), center_cell_style)
            table_data[16][14] = Paragraph(str(overhaul_data[28][34]), center_cell_style)
            table_data[16][15] = Paragraph(str(overhaul_data[28][36]), center_cell_style)
            table_data[16][16] = Paragraph(str(overhaul_data[28][38]), center_cell_style)

            table_data[17][0] = Paragraph(str(overhaul_data[29][0]), cell_style_bold)
            table_data[17][7] = Paragraph(str(overhaul_data[29][17]), center_cell_style_bold)
            table_data[17][8] = Paragraph(str(overhaul_data[29][19]), center_cell_style_bold)
            table_data[17][9] = Paragraph(str(overhaul_data[29][21]), center_cell_style_bold)
            table_data[17][10] = Paragraph(str(overhaul_data[29][24]), center_cell_style_bold)
            table_data[17][11] = Paragraph(str(overhaul_data[29][27]), center_cell_style_bold)
            table_data[17][12] = Paragraph(str(overhaul_data[29][30]), center_cell_style_bold)
            table_data[17][13] = Paragraph(str(overhaul_data[29][32]), center_cell_style_bold)
            table_data[17][14] = Paragraph(str(overhaul_data[29][34]), center_cell_style_bold)

            table_data[18][0] = Paragraph(str(overhaul_data[30][0]), center_cell_style)

            table = Table(table_data, style=[
                ('SPAN', (0, 0), (-1, 0)),
                ('SPAN', (0, 1), (1, 1)),
                ('SPAN', (2, 1), (6, 1)),
                ('SPAN', (9, 1), (14, 1)),
                ('SPAN', (15, 1), (16, 9)),
                ('SPAN', (0, 2), (1, 2)),
                ('SPAN', (2, 2), (6, 2)),
                ('SPAN', (9, 2), (11, 2)),
                ('SPAN', (0, 3), (6, 3)),
                ('SPAN', (9, 3), (11, 3)),
                ('SPAN', (0, 4), (6, 4)),
                ('SPAN', (8, 4), (11, 4)),
                ('SPAN', (0, 5), (1, 5)),
                ('SPAN', (2, 5), (6, 5)),
                ('SPAN', (8, 5), (10, 5)),
                ('SPAN', (0, 6), (1, 6)),
                ('SPAN', (2, 6), (3, 6)),
                ('SPAN', (4, 6), (5, 6)),
                ('SPAN', (8, 6), (10, 6)),
                ('SPAN', (0, 7), (1, 7)),
                ('SPAN', (2, 7), (3, 7)),
                ('SPAN', (4, 7), (5, 7)),
                ('SPAN', (8, 7), (10, 7)),
                ('SPAN', (0, 8), (1, 8)),
                ('SPAN', (2, 8), (6, 8)),
                ('SPAN', (8, 8), (10, 8)),
                ('SPAN', (8, 9), (10, 9)),
                ('SPAN', (0, 11), (1, 11)),
                ('SPAN', (2, 11), (-1, 11)),
                ('SPAN', (0, 12), (0, 14)),
                ('SPAN', (1, 12), (1, 14)),
                ('SPAN', (2, 12), (5, 13)),
                ('SPAN', (6, 12), (6, 14)),
                ('SPAN', (7, 12), (8, 13)),
                ('SPAN', (9, 12), (9, 14)),
                ('SPAN', (10, 12), (10, 14)),
                ('SPAN', (11, 12), (11, 14)),
                ('SPAN', (12, 12), (12, 14)),
                ('SPAN', (13, 12), (13, 14)),
                ('SPAN', (14, 12), (16, 12)),
                ('SPAN', (14, 13), (14, 14)),
                ('SPAN', (15, 13), (16, 13)),
                ('SPAN', (2, 15), (3, 15)),
                ('SPAN', (4, 15), (5, 15)),
                ('SPAN', (0, 17), (6, 17)),
                ('SPAN', (0, 18), (-1, 18)),
                ('SPAN', (12, 4), (14, 9)),
                ('LEFTPADDING', (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
                ('GRID', (0, 1), (6, 8), 0.5, colors.black),
                ('GRID', (8, 1), (14, 9), 0.5, colors.black),
                ('GRID', (0, 11), (-1, 16), 0.5, colors.black),
                ('GRID', (0, 17), (14, 17), 0.5, colors.black),
                ('BOX', (0, 0), (-1, -1), 1, colors.black),
                ('BOX', (0, 0), (-1, 0), 1, colors.black),
                ('BOX', (0, 1), (6, 8), 1, colors.black),
                ('BOX', (8, 1), (14, 9), 1, colors.black),
                ('BOX', (0, 11), (-1, 17), 1, colors.black),
                ('BOX', (15, 17), (16, 17), 1, colors.black),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke),
                ('BACKGROUND', (0, -1), (-1, -1), colors.whitesmoke),
                ('BACKGROUND', (7, 1), (7, 8), colors.lightgrey),
                ('BACKGROUND', (0, 9), (7, 10), colors.lightgrey),
                ('BACKGROUND', (8, 10), (-1, 10), colors.lightgrey),
                ('BACKGROUND', (15, 1), (16, 10), colors.lightgrey),
                ('BACKGROUND', (15, 17), (16, 17), colors.lightgrey),
            ], colWidths=[60, 40, 40, 35, 40, 35, 45, 50, 50, 55, 55, 65, 65, 50, 40, 40, 45])
            content.append(table)
            content.append(PageBreak())

        pdf.build(content)

        wb.close()
        return pdf_file_path

    except Exception as e:
        error_msg = str(e)
        print(f"Error in overhaul_to_pdf: {error_msg}", file=sys.stderr)
        messagebox.showerror("Error", error_msg)
        return None


def center_window(window):
    window.update_idletasks()
    width = window.winfo_width()
    height = window.winfo_height()
    x = (window.winfo_screenwidth() // 2) - (width // 2)
    y = (window.winfo_screenheight() // 2) - (height // 2)
    window.geometry(f'{width}x{height}+{x}+{y}')


def open_overhaul():
    file_path = filedialog.askopenfilename(filetypes=[('Excel file', '*.xlsx')])
    if file_path:
        # Создаем окно загрузки
        loading_window = tk.Toplevel(root)
        loading_window.title("Обработка...")
        loading_window.geometry("300x100")
        loading_window.resizable(False, False)
        
        # Центрируем окно по экрану
        center_window(loading_window)
        loading_window.transient(root)
        loading_window.grab_set()
        
        label = tk.Label(loading_window, text="Работаю с файлом . . .", font=('Times New Roman', 10))
        label.pack(expand=True)
        
        # Обновляем GUI перед обработкой
        root.update()
        
        pdf_file_path = overhaul_to_pdf(file_path)
        
        # Закрываем окно загрузки
        loading_window.destroy()
        
        if pdf_file_path:
            messagebox.showinfo("Success", f"PDF создан: {pdf_file_path}")


# Создание GUI
root = tk.Tk()
root.title("Receipt Converter")
root.geometry("300x150")
center_window(root)

# Кнопка для выбора файла
overhaul_button = tk.Button(root, text="Капремонт", command=open_overhaul)
overhaul_button.pack(pady=20)
receipt_button = tk.Button(root, text="Коммунальные платежи", command=open_receipt)
receipt_button.pack(pady=20)

# Запуск GUI
root.mainloop()